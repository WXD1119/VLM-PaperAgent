from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 100, 112)


def font(run, name="Microsoft YaHei", size=11, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}")) or OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            tc_mar.append(node)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, "F2F4F7")
        for run in cell.paragraphs[0].runs:
            font(run, size=9.5, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            for run in cells[i].paragraphs[0].runs:
                font(run, size=9.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def setup_doc(title, subtitle, doc_type):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run(f"VLM-PaperAgent | {doc_type}"), size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("v1.2 · 2026-07"), size=8.5, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(title), size=23, bold=True, color=RGBColor(0, 0, 0))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    font(p.add_run(subtitle), size=12, color=MUTED)
    add_table(doc, ["版本", "定位", "状态"], [["v1.2", "证据可追溯 + 可评测 + 可恢复", "代码骨架已建立"]], [1200, 5360, 2800])
    return doc


def h(doc, text, level=1):
    return doc.add_heading(text, level=level)


def p(doc, text, bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        font(para.add_run(bold_prefix), bold=True)
        font(para.add_run(text[len(bold_prefix):]))
    else:
        font(para.add_run(text))
    return para


def bullets(doc, items: Iterable[str]):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.left_indent = Inches(0.5)
        para.paragraph_format.first_line_indent = Inches(-0.25)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.line_spacing = 1.167
        font(para.add_run(item))


def save(doc, filename):
    path = ROOT / filename
    doc.save(path)
    print(path.name)


def build_prd():
    doc = setup_doc("多模态文献精读智能体系统 PRD", "求职项目落地版：核心价值与技术栈并重", "Product Requirements")
    h(doc, "1. 产品定位")
    p(doc, "面向 CV/VLM 论文的证据可追溯精读 Agent。系统将文本、公式、表格和图像说明统一建模，使问答、总结与审稿疑点均能回溯到页码、章节和证据元素。")
    h(doc, "1.1 核心用户价值", 2)
    bullets(doc, ["降低论文中公式、表格和跨章节信息的查找成本。", "输出不是无来源的结论，而是带证据定位和置信度的结构化结果。", "对论文主张给出 supported、questionable 或 insufficient_evidence 判定，辅助人工复核。"])
    h(doc, "2. MVP 边界")
    add_table(doc, ["层级", "能力", "技术栈", "验收标准"], [
        ["核心 MVP", "解析、索引、问答、审稿、评测", "Marker/MinerU、Chroma、BM25、Reranker、Pydantic、Streamlit", "单篇真实论文无 Mock 跑通"],
        ["简历增强", "API、会话、图谱、可观测", "FastAPI、Redis、Neo4j、Tracing", "10 篇论文跨文献演示"],
        ["二期扩展", "异步批处理与周报", "Celery、消息队列、PostgreSQL", "任务幂等、失败重试"],
    ], [1200, 2800, 3160, 2200])
    h(doc, "3. 端到端用户流程")
    bullets(doc, ["上传 PDF，生成 paper_id 与内容哈希。", "解析文本、公式、表格、图片说明，并保留 page、bbox、section_path。", "Dense 与 BM25 双路召回，经 RRF 和 CrossEncoder 重排。", "Claim Extractor 抽取主张，Evidence Agent 检索证据，Critic/Judge 完成审稿。", "Reporter 输出含 evidence_ids 的报告，界面可跳转到原页。"])
    h(doc, "4. 成功指标")
    add_table(doc, ["层级", "指标", "口径"], [
        ["解析", "元素识别 F1、页码定位准确率", "按文本/公式/表格/图注分别统计"],
        ["检索", "Recall@5、MRR、nDCG@10", "Golden Set 人工标注证据 ID"],
        ["生成", "Citation Precision/Completeness", "答案陈述是否有正确证据支撑"],
        ["审稿", "Precision、Recall、F1", "疑点与人工标注的一致性"],
    ], [1400, 2800, 5160])
    h(doc, "5. 四周里程碑")
    add_table(doc, ["周次", "交付物", "Demo 检查点"], [
        ["第1周", "领域模型、Parser、父子切片", "真实论文结构化结果可查看"],
        ["第2周", "Dense+BM25+RRF+Rerank", "问答带页码证据；完成消融评测"],
        ["第3周", "手写 FSM、Reviewer、checkpoint", "失败可重试并从 checkpoint 恢复"],
        ["第4周", "FastAPI、Streamlit、Golden Set、CI", "完整 Demo 与真实指标"],
    ], [1000, 3900, 4460])
    h(doc, "6. 非目标与诚实边界")
    bullets(doc, ["不宣称自动证明论文错误，只输出值得人工复核的证据化疑点。", "核心 MVP 不依赖 Neo4j、Redis 或 Celery 才能运行。", "简历数字必须来自固定测试集，不使用预设宣传指标。"])
    save(doc, "多模态文献精读智能体系统产品需求文档_v1.2.docx")


def build_fsd():
    doc = setup_doc("VLM-PaperAgent 功能设计说明书", "模块职责、数据流与手写 FSM 设计", "Functional Design")
    h(doc, "1. 总体架构")
    p(doc, "系统采用领域层、应用工作流层、基础设施层与表现层分离。Agent 是执行特定业务职责的节点，不直接访问 UI，也不能任意修改状态。")
    add_table(doc, ["层", "目录", "职责"], [
        ["领域层", "domain/", "PaperElement、Claim、Evidence、ReviewFinding、RunContext"],
        ["应用层", "workflow/、agents/", "节点编排、状态转换、重试、报告生成"],
        ["基础设施", "ingestion/、retrieval/、storage/、llm/", "解析、检索、数据库与模型适配"],
        ["表现层", "api/、app/", "FastAPI 服务与 Streamlit Demo"],
    ], [1400, 2500, 5460])
    h(doc, "2. 数据摄取与证据模型")
    bullets(doc, ["Parser 通过统一 PaperParser Protocol 封装 Marker/MinerU，避免业务层绑定具体工具。", "每个元素保存 paper_id、page、bbox、section_path、parent/prev/next 关系和 parser_version。", "采用父子切片：子元素负责检索，父章节负责补充生成上下文。", "公式、表格、图注与前后解释段落建立显式关联。"])
    h(doc, "3. 混合检索")
    p(doc, "Dense Top-20 与 BM25 Top-20 分别召回，经 Reciprocal Rank Fusion 统一排序，再由 CrossEncoder 重排至 Top-8，最后进行邻接上下文扩展。")
    h(doc, "4. Reviewer 工作流")
    add_table(doc, ["节点", "输入", "输出", "失败策略"], [
        ["Claim Extractor", "结构化论文", "Claim[]", "结构化校验失败重试"],
        ["Evidence Agent", "Claim+检索器", "Evidence IDs", "无证据转 insufficient"],
        ["Critic", "Claim+Evidence", "批判摘要", "超时重试"],
        ["Judge", "证据与批判摘要", "结构化 verdict", "低置信度转人工复核"],
        ["Reflector", "冲突上下文", "修订后的判断", "达到上限转人工复核"],
        ["Reporter", "ReviewFinding[]", "ReviewReport", "禁止缺失证据的 supported"],
    ], [1700, 2500, 2500, 2660])
    h(doc, "5. 手写 FSM 调度器")
    p(doc, "主路径：INIT → PARSING → INDEXING → EXTRACTING_CLAIMS → RETRIEVING_EVIDENCE → CRITIQUING → JUDGING → REPORTING → FINISHED。Judge 可转 REFLECTING 或 HUMAN_REVIEW。")
    bullets(doc, ["合法转换集中定义；非法转换立即失败。", "节点接收 RunContext 副本，只返回 NodeResult 增量。", "节点级 attempt 计数、最大步骤数和异常隔离。", "每次状态转换保存 checkpoint，支持重启恢复。", "外部写入必须以 run_id + node + artifact_type 设计幂等键。"])
    h(doc, "6. 可观测性与安全")
    bullets(doc, ["记录 run_id、paper_id、node、attempt、latency、model_version、prompt_version。", "上传文件使用内容哈希命名，限制大小和 MIME 类型，禁止直接使用原文件名写路径。", "日志不保存完整思维链，仅保存证据、判定摘要和结构化错误。", "模型输出必须经 Pydantic 校验，失败结果不得污染后续状态。"])
    h(doc, "7. 仓库映射")
    p(doc, "代码骨架已按 domain、ingestion、retrieval、agents、workflow、storage、llm、evaluation、observability、api 分包；测试覆盖 RRF、证据约束与 FSM Happy Path。")
    save(doc, "功能设计文档FSD_v1.2.docx")


def build_api():
    doc = setup_doc("VLM-PaperAgent 接口文档", "内部节点契约、基础设施端口与 HTTP API", "API Specification")
    h(doc, "1. 统一原则")
    bullets(doc, ["本版本以手写 FSM 为唯一主调度实现，LangGraph 仅作为二期对照方案。", "跨层交互使用 Pydantic 模型或 Protocol，不传递无约束 Dict。", "所有审稿结论必须返回 evidence_ids；所有运行记录携带 run_id。"])
    h(doc, "2. 核心内部契约")
    add_table(doc, ["接口", "签名", "说明"], [
        ["PaperParser", "parse(pdf_path, paper_id) -> Paper", "解析器适配端口"],
        ["WorkflowNode", "execute(context) -> NodeResult", "工作流节点统一接口"],
        ["CheckpointStore", "save(context); load(run_id)", "执行状态持久化"],
        ["VectorStore", "upsert(elements); search(query, top_k)", "向量存储端口"],
        ["LLMClient", "generate_structured(prompt, model)", "结构化模型输出"],
    ], [1700, 3900, 3760])
    h(doc, "3. NodeResult 契约")
    p(doc, "字段：next_state 表示请求的目标状态；artifacts 表示节点产出的结构化增量；message 表示面向日志的简短摘要。Scheduler 在合并 artifacts 前验证状态转换。")
    h(doc, "4. HTTP API 草案")
    add_table(doc, ["方法", "路径", "用途", "响应"], [
        ["GET", "/health", "健康检查", "{status: ok}"],
        ["POST", "/papers", "上传论文并创建 paper_id", "PaperSummary"],
        ["POST", "/runs", "启动解析/审稿工作流", "RunStatus"],
        ["GET", "/runs/{run_id}", "查询状态、轨迹与错误", "RunDetail"],
        ["POST", "/papers/{paper_id}/questions", "带证据问答", "AnswerWithCitations"],
        ["GET", "/papers/{paper_id}/report", "获取审稿报告", "ReviewReport"],
    ], [1000, 3000, 2900, 2460])
    h(doc, "5. 错误模型")
    p(doc, "统一错误字段包括 code、message、run_id、state、retryable 和 details。模型超时、输出校验失败、解析失败与基础设施不可用必须使用不同 code。")
    h(doc, "6. 版本与幂等")
    bullets(doc, ["Prompt、Parser、Embedding 和 Reranker 均记录版本。", "POST /runs 接受 Idempotency-Key，重复请求返回同一 run。", "数据库写入使用 paper_id 或 run_id 作为天然分区键。"])
    save(doc, "接口文档_v1.2.docx")


def build_dictionary():
    doc = setup_doc("VLM-PaperAgent 数据字典", "领域模型与 Chroma / Redis / Neo4j Schema", "Data Dictionary")
    h(doc, "1. 领域模型")
    add_table(doc, ["实体", "关键字段", "约束"], [
        ["Paper", "paper_id, title, source_path, sha256", "sha256 用于文件去重"],
        ["PaperElement", "element_id, page, bbox, section_path, element_type, content", "page>=1；content 非空"],
        ["Claim", "claim_id, paper_id, text, source_element_ids", "必须可定位原文主张"],
        ["ReviewFinding", "claim_id, verdict, evidence_ids, confidence", "supported 必须有 evidence_ids"],
        ["RunContext", "run_id, state, artifacts, errors, state_history", "每次转换持久化"],
    ], [1600, 5000, 2760])
    h(doc, "2. Chroma：paper_elements")
    add_table(doc, ["字段", "位置", "说明"], [
        ["element_id", "document id", "全局唯一"],
        ["content", "document", "用于 Embedding 的规范化文本"],
        ["paper_id/page", "metadata", "论文过滤与引用定位"],
        ["element_type", "metadata", "paragraph/equation/table/figure/caption"],
        ["section_path", "metadata JSON", "章节层级"],
        ["parent_id/prev_id/next_id", "metadata", "上下文扩展"],
        ["embedding_version", "metadata", "支持重建索引与回归"],
    ], [2200, 2200, 4960])
    h(doc, "3. Redis")
    add_table(doc, ["Key 模式", "数据", "TTL/用途"], [
        ["paper-agent:run:{run_id}", "RunContext JSON", "7天；checkpoint"],
        ["paper-agent:session:{session_id}:history", "消息列表", "24小时；短期会话"],
        ["paper-agent:idempotency:{key}", "run_id", "24小时；防止重复任务"],
        ["paper-agent:lock:{paper_id}", "owner token", "短 TTL；防止重复索引"],
    ], [3600, 3000, 2760])
    h(doc, "4. Neo4j")
    add_table(doc, ["类型", "Label/Relationship", "关键属性"], [
        ["节点", "Paper", "paper_id, title, year, arxiv_id"],
        ["节点", "Method", "canonical_name, aliases, embedding"],
        ["节点", "Dataset", "canonical_name, task"],
        ["节点", "MetricResult", "metric, value, split"],
        ["节点", "Insight", "verdict, confidence, evidence_ids"],
        ["关系", "PROPOSES / EVALUATED_ON / IMPROVES_UPON", "reason, evidence_ids, run_id"],
    ], [1200, 3800, 4360])
    h(doc, "5. 运行与评测记录")
    add_table(doc, ["实体", "关键字段", "用途"], [
        ["RunRecord", "run_id, node, attempt, latency_ms, status", "可观测与失败定位"],
        ["ModelCall", "model_version, prompt_version, tokens, latency", "成本与回归"],
        ["GoldenQuestion", "question_id, paper_id, relevant_element_ids", "检索评测"],
        ["EvalResult", "dataset_version, metric, value, commit_sha", "防止指标口径漂移"],
    ], [1800, 4400, 3160])
    h(doc, "6. 一致性规则")
    bullets(doc, ["paper_id 与 element_id 在所有存储中保持一致。", "Chroma 为检索索引，不作为业务事实的唯一来源。", "Neo4j 关系必须携带 evidence_ids，避免无来源图谱边。", "Redis checkpoint 只保存可序列化状态，不保存数据库连接或模型客户端。"])
    save(doc, "数据字典_v1.2.docx")


if __name__ == "__main__":
    build_prd()
    build_fsd()
    build_api()
    build_dictionary()
